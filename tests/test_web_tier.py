"""Web-verified tier: fully cited from public pages, but not confirmed on official registries."""
from datetime import date

from app import appsettings, excel, report
from app import rules as R
from app.ingest import EMPTY_STATS, ingest_candidate, refresh_web_tier
from app.models import Lead, PipelineEntry, Run, utcnow

from .conftest import GOOD
from .test_web import client, csrf, login  # noqa: F401  (client is a fixture)

# SYNTHETIC: same fake company, but the registries could not be opened - only directory / own-site pages.
WEB = {**GOOD, "verification_status": "PARTIALLY VERIFIED",
       "indian_registry_url": "https://www.companydirectory.example/testco-cin",
       "us_registry_url": "", "us_verification_limitation": "wyobiz.wyo.gov requires a CAPTCHA"}


def test_cited_online_is_web_verified_but_not_verified(db):
    st = dict(EMPTY_STATS)
    res = ingest_candidate(db, WEB, None, st)
    lead = db.query(Lead).one()
    assert R.gate(WEB)[0] != "VERIFIED" and lead.verification_status != "VERIFIED"  # strict gate unchanged
    assert res["web_verified"] and lead.web_verified_at and st["web_verified"] == 1 and st["new_verified"] == 0
    assert db.query(PipelineEntry).count() == 1


def test_missing_citation_is_not_web_verified(db):
    ok, missing = R.web_evidence({**WEB, "cin_llpin": ""})
    assert not ok and "Indian legal name with a valid CIN/LLPIN" in missing
    ok, missing = R.web_evidence({**WEB, "relationship_evidence_urls": []})
    assert not ok and "a cited page linking the two entities" in missing
    ingest_candidate(db, {**WEB, "relationship_evidence": ""})
    assert db.query(Lead).one().web_verified_at is None


def test_refresh_backfills_existing_queue(db):
    ingest_candidate(db, WEB)
    lead = db.query(Lead).one()
    lead.web_verified_at = None
    db.commit()
    assert refresh_web_tier(db) == 1 and db.query(Lead).one().web_verified_at


def test_lists_report_and_workbook(client, db):
    from io import BytesIO

    from openpyxl import load_workbook
    db.query(Lead).delete()
    db.query(PipelineEntry).delete()
    db.commit()
    run = Run(run_date=date.today(), status="completed", stats={"web_verified": 1}, started_at=utcnow())
    db.add(run)
    db.commit()
    ingest_candidate(db, WEB)
    run.finished_at = utcnow()
    db.commit()
    login(client, "admin@example.com")
    page = client.get("/leads?status=web").text
    assert "Web-verified leads" in page and "Testco" in page
    assert "Testco" not in client.get("/leads?status=queue").text
    assert "1 web-verified" in client.get("/").text
    subject, text, html_body = report.build(db, run, appsettings.get_all(db))
    assert "+1 web-verified" in subject and "New web-verified leads" in html_body and "Testco" in text
    wb = load_workbook(BytesIO(excel.build_workbook(db)))
    assert wb["Web-Verified Leads"].max_row == 2 and wb["Research Queue"].max_row == 1


def test_running_run_shows_finds_live(client, db):
    run = Run(run_date=date.today(), trigger="manual", status="running", stats={}, started_at=utcnow())
    db.add(run)
    db.commit()
    ingest_candidate(db, {**WEB, "cin_llpin": "U72900KA2019PTC777777", "us_entity_number": "77",
                          "indian_legal_name": "Livetest Private Limited", "indian_brand_name": "Livetest",
                          "indian_website": "https://live.example"}, run.id)
    login(client, "admin@example.com")
    d = client.get(f"/api/runs/{run.id}/status").json()
    assert d["stats"]["candidates"] == 1 and d["stats"]["web_verified"] == 1
    assert d["recent"][0]["name"] == "Livetest" and d["recent"][0]["status"] == "Web-verified"
    assert "Found in this run" in client.get(f"/runs/{run.id}").text


def test_workbook_has_only_the_kept_columns_and_unused_fields_are_not_collected(db):
    from io import BytesIO

    from openpyxl import load_workbook

    from app.fields import EXCEL_FIELDS
    ingest_candidate(db, WEB)
    lead = db.query(Lead).one()
    assert lead.recent_activity_evidence is None and lead.linkedin_decision_maker_url is None  # not researched
    assert lead.cin_llpin  # still kept internally as evidence for the tiers
    wb = load_workbook(BytesIO(excel.build_workbook(db)))
    assert "Verified Leads" not in wb.sheetnames
    for sheet in ("Web-Verified Leads", "Research Queue"):
        headers = [c.value for c in wb[sheet][1]]
        assert headers == [label for _, label, _ in EXCEL_FIELDS] and len(headers) == 18
        assert "CIN/LLPIN" not in headers and "Lead ID" not in headers
