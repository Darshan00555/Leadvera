"""Exports the database as the master workbook (Section 8 sheets plus web-verified leads)."""
import io

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from .fields import EXCEL_FIELDS, LABEL, QUEUE_STATUSES
from .models import Coverage, Exclusion, Lead, PipelineEntry, Run

HEADER_FILL = PatternFill("solid", fgColor="14284B")
HEADER_FONT = Font(bold=True, color="FFFFFF")
WIDE_HINTS = ("Description", "Evidence", "URLs", "Angle", "Notes", "Limitations", "Missing", "Signals",
              "Reason", "Recommendations", "Errors", "Sources", "Response", "Outcome")


def _sheet(wb, title, headers, rows, freeze="B2"):
    ws = wb.create_sheet(title)
    ws.append(headers)
    for i, h in enumerate(headers, 1):
        c = ws.cell(row=1, column=i)
        c.fill, c.font = HEADER_FILL, HEADER_FONT
        c.alignment = Alignment(wrap_text=True, vertical="center")
        ws.column_dimensions[get_column_letter(i)].width = 48 if any(w in h for w in WIDE_HINTS) else \
            max(12, min(30, len(h) + 4))
    for r in rows:
        ws.append([_cell(v) for v in r])
    for row in ws.iter_rows(min_row=2):
        for c in row:
            if hasattr(c.value, "year"):
                c.number_format = "yyyy-mm-dd"
            c.alignment = Alignment(vertical="top", wrap_text=isinstance(c.value, str) and len(c.value) > 60)
    ws.freeze_panes = freeze
    ws.auto_filter.ref = f"A1:{get_column_letter(len(headers))}{max(1, ws.max_row)}"
    ws.row_dimensions[1].height = 32
    return ws


def _cell(v):
    if isinstance(v, (list, tuple)):
        return "\n".join(map(str, v))
    if isinstance(v, dict):
        return ", ".join(f"{k}={x}" for k, x in v.items())
    return v


def build_workbook(db):
    wb = Workbook()
    wb.remove(wb.active)
    lead_attrs = [a for a, _, _ in EXCEL_FIELDS]

    # No separate "Verified Leads" sheet: a registry-verified lead is listed with the web-verified ones.
    verified = db.query(Lead).filter_by(verification_status="VERIFIED").order_by(Lead.id).all()
    queue = db.query(Lead).filter(Lead.verification_status.in_(QUEUE_STATUSES)).order_by(Lead.id).all()
    web = verified + [l for l in queue if l.web_verified_at]
    _sheet(wb, "Web-Verified Leads", [LABEL[a] for a in lead_attrs],
           [[getattr(l, a) for a in lead_attrs] for l in web])
    _sheet(wb, "Research Queue", [LABEL[a] for a in lead_attrs],
           [[getattr(l, a) for a in lead_attrs] for l in queue if not l.web_verified_at])

    runs = db.query(Run).order_by(Run.id).all()
    _sheet(wb, "Daily Run Log",
           ["Run ID", "Run Date", "Trigger", "Status", "Started At", "Finished At", "Processing Minutes",
            "Queries Executed", "Query Groups", "Sources Consulted", "Candidates Found", "Indian Verified",
            "US LLC Verified", "Relationships Verified", "New Verified Leads", "Records Updated",
            "Partially Verified", "Duplicates", "Disqualified", "Errors", "Blocked Sources", "Stop Reason",
            "Cost USD", "Recommendations", "Report Sent At"],
           [[r.id, r.run_date, r.trigger, r.status, r.started_at, r.finished_at,
             round((r.finished_at - r.started_at).total_seconds() / 60, 1) if r.started_at and r.finished_at else None,
             len(r.queries or []), r.query_groups, r.sources, (r.stats or {}).get("candidates"),
             (r.stats or {}).get("indian_verified"), (r.stats or {}).get("us_verified"),
             (r.stats or {}).get("relationships_verified"), (r.stats or {}).get("new_verified"),
             (r.stats or {}).get("updated"), (r.stats or {}).get("partial"), (r.stats or {}).get("duplicates"),
             (r.stats or {}).get("disqualified"), r.errors, r.blocked_sources, (r.stats or {}).get("stop_reason"),
             r.cost_usd, r.recommendations, r.report_sent_at] for r in runs], freeze="A2")

    ex = db.query(Exclusion).order_by(Exclusion.id).all()
    _sheet(wb, "Exclusions and Duplicates",
           ["Exclusion ID", "Date", "Candidate Name", "CIN/LLPIN", "US State", "US Entity Number", "Domain",
            "Type", "Reason", "Canonical Lead ID", "Source URLs", "Resolved"],
           [[f"EX-{x.id:06d}", x.date, x.candidate_name, x.cin_llpin, x.us_state, x.us_entity_number, x.domain,
             x.type, x.reason, x.canonical_lead_id, x.source_urls, "Yes" if x.resolved else "No"] for x in ex],
           freeze="A2")

    pipe = db.query(PipelineEntry, Lead).join(Lead, Lead.id == PipelineEntry.lead_pk).order_by(PipelineEntry.id).all()
    _sheet(wb, "Sales Pipeline",
           ["Lead ID", "Company", "Priority", "Lead Owner", "Outreach Date", "Channel", "Last Contact", "Response",
            "Meeting Booked", "Opportunity Stage", "Next Follow-Up Date", "Est. Monthly Recurring Revenue (USD)",
            "Actual Outcome", "Notes"],
           [[l.lead_id, l.indian_brand_name or l.indian_legal_name, l.priority, p.lead_owner, p.outreach_date,
             p.channel, p.last_contact, p.response, "Yes" if p.meeting_booked else "No", p.opportunity_stage,
             p.next_follow_up, p.est_mrr_usd, p.actual_outcome, p.notes] for p, l in pipe])

    cov = db.query(Coverage).order_by(Coverage.id).all()
    _sheet(wb, "Research Coverage",
           ["Date", "Query Group", "Industry", "Indian State", "US State", "Pattern", "Sources", "Queries Run",
            "Candidates Yielded", "Verified Yielded", "Run ID"],
           [[c.date, c.query_group, c.industry, c.indian_state, c.us_state, c.pattern, c.sources, c.queries_run,
             c.candidates_yielded, c.verified_yielded, c.run_id] for c in cov], freeze="A2")

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
